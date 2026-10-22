"""Offline tier: replay a case's recorded trace through today's formatting code."""

from __future__ import annotations

import csv
import io
import json
import os
import sys

APP = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(APP, "src"))

from attachments import build_file, citable_tables, file_format  # noqa: E402
from slack_blocks import parse_markdown_table, tool_leaf  # noqa: E402
from slack_format import format_for_slack, split_markdown_tables  # noqa: E402
from slack_search import TOOL_LIST_CHANNEL  # noqa: E402
from spreadsheet import SheetTable, tables_from_steps  # noqa: E402

from evals import checks  # noqa: E402
from evals.case import CaseResult  # noqa: E402


def _listing_tables(steps: list) -> list:
    """Slack listing tables as recorded (the live app attaches the full Listing; a trace keeps its first rows)."""
    tables = []
    for step in steps:
        if tool_leaf(step.get("tool_name") or "") != TOOL_LIST_CHANNEL or step.get("status") != "success":
            continue
        output = str(step.get("output") or "")
        title = next((ln.lstrip("#").strip() for ln in output.splitlines() if ln.startswith("#")), "Slack listing")
        for kind, chunk in split_markdown_tables(output):
            headers, rows = parse_markdown_table(chunk) if kind == "table" else ([], [])
            if headers and rows:
                tables.append(SheetTable(title, headers, rows))
    return tables


def _attachment_reasons(case: dict, steps: list) -> list[str]:
    fmt = file_format(case["question"])
    if not fmt:
        return ["expected attachment, but the question doesn't ask for Excel / CSV / JSON / RIS"]
    tables = _listing_tables(steps) + tables_from_steps(steps)
    if not tables:
        return ["expected attachment, but no table-producing tool output"]
    if fmt == "ris" and not citable_tables(tables):
        return ["asked for RIS, but no tool produced links or files to cite"]
    content, filename, _ = build_file(fmt, tables)
    if not filename.endswith("." + fmt):
        return [f"asked for {fmt}, built {filename}"]
    if fmt == "json":
        doc = json.loads(content)
        docs = doc if isinstance(doc, list) else [doc]
        if len(docs) != len(tables) or not all(d.get("rows") for d in docs):
            return [f"JSON has {len(docs)} tables with rows for {len(tables)} tables"]
    elif fmt == "ris":
        text = content.decode("utf-8")
        records = [r for r in text.split("ER  - ") if r.strip()]
        rows = sum(len(t.rows) for t in citable_tables(tables))
        if len(records) != rows or not all(r.strip().startswith("TY  - ") for r in records):
            return [f"RIS has {len(records)} well-formed records for {rows} rows"]
    elif fmt == "csv":
        if len(list(csv.reader(io.StringIO(content.decode("utf-8-sig"))))) < 2:
            return ["CSV has no data rows"]
    else:
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content))
        if len(wb.sheetnames) != len(tables):
            return [f"workbook has {len(wb.sheetnames)} sheets for {len(tables)} tables"]
    return []


def run_offline(case: dict) -> CaseResult:
    rec = case.get("recorded")
    if not rec:
        return CaseResult(case["id"], True, [], skipped=True)
    text = ""
    try:
        steps = rec["steps"]
        text, blocks = format_for_slack(rec["final_text"], use_llm=False, tool_steps=steps)
        expect = case.get("expect") or {}
        reasons = checks.universal(text, blocks)
        if not expect.get("allow_condensed"):
            reasons += checks.not_collapsed(rec["final_text"], blocks)
            reasons += checks.structure_kept(rec["final_text"], blocks)
        reasons += checks.blocks_present(blocks, expect.get("blocks") or [])
        # Only what Slack shows counts: the fallback text always holds the full answer.
        shown = checks.visible_text(blocks)
        reasons += checks.phrases_all(shown, expect.get("phrases") or [])
        reasons += checks.forbid(shown, expect.get("forbid") or [])
        if expect.get("attachment"):
            reasons += _attachment_reasons(case, steps)
    except Exception as e:  # a broken case or replay must not stop the run
        reasons = [f"exception: {type(e).__name__}: {e}"]
    return CaseResult(case["id"], not reasons, reasons, slack_text=text)
