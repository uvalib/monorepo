"""
spreadsheet.py — Excel (.xlsx) attachments for occupancy answers.

Built straight from the occupancy tools' markdown tables (never from the
model's prose), so numbers in the file match the tool output exactly.
app.py attaches the workbook when the user's message asks for Excel /
a spreadsheet / CSV and a table-producing tool ran this turn.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Any, Iterable, List, Optional

from slack_blocks import parse_markdown_table, tool_leaf
from slack_format import split_markdown_tables

# Tools whose output tables are worth exporting (hours/spaces are not).
TABLE_TOOLS = frozenset(
    {"get_traffic_by_period", "get_entrance_traffic", "get_occupancy_report"}
)

_WANTS = re.compile(r"(?i)\b(?:excel|spreadsheets?|xlsx|csv)\b")
_NUMBER = re.compile(r"^-?[\d,]+(?:\.\d+)?$")
_PERCENT = re.compile(r"^-?\d+(?:\.\d+)?%$")
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_SHEET_BAD = re.compile(r"[\[\]:*?/\\]")


@dataclass
class SheetTable:
    title: str
    headers: List[str]
    rows: List[List[str]] = field(default_factory=list)
    # True for people's text (Slack listings): cells stay exact strings, never numbers or formulas.
    plain: bool = False
    source: str = ""  # where the rows came from, e.g. "#ai" (RIS notes)


def wants_spreadsheet(user_text: str) -> bool:
    return bool(_WANTS.search(user_text or ""))


def _step_attr(step: Any, key: str) -> Any:
    if isinstance(step, dict):
        return step.get(key)
    return getattr(step, key, None)


def tables_from_steps(steps: Optional[Iterable[Any]]) -> List[SheetTable]:
    tables: List[SheetTable] = []
    for step in steps or []:
        if _step_attr(step, "status") not in (None, "success"):
            continue
        if tool_leaf(_step_attr(step, "tool_name") or "") not in TABLE_TOOLS:
            continue
        output = str(_step_attr(step, "output") or "")
        title = ""
        for line in output.splitlines():
            if line.startswith("#"):
                title = line.lstrip("#").strip()
                break
        for kind, chunk in split_markdown_tables(output):
            if kind != "table":
                continue
            headers, rows = parse_markdown_table(chunk)
            if headers and rows:
                tables.append(SheetTable(title or "HooHelp data", headers, rows))
    return tables


def _cell_value(raw: str):
    s = (raw or "").strip().strip("*`")
    if _PERCENT.match(s):
        return float(s[:-1]) / 100, "0.0%"
    if _NUMBER.match(s):
        num = float(s.replace(",", ""))
        return (int(num) if num.is_integer() else num), "#,##0" if num.is_integer() else "#,##0.00"
    if _DATE.match(s):
        return datetime.strptime(s, "%Y-%m-%d"), "yyyy-mm-dd"
    return s, None


def _sheet_name(title: str, used: set) -> str:
    base = _SHEET_BAD.sub("", title.split("(")[0]).strip() or "Data"
    name = base[:31]
    n = 2
    while name in used:
        suffix = f" ({n})"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name)
    return name


def _plain_cell(raw: str) -> str:
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    return ILLEGAL_CHARACTERS_RE.sub("", str(raw or ""))


def build_xlsx(tables: List[SheetTable]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    for table in tables:
        ws = wb.create_sheet(_sheet_name(table.title, used))
        ws.append([table.title])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([])
        ws.append(table.headers)
        for cell in ws[3]:
            cell.font = Font(bold=True)
        widths = [len(h) for h in table.headers]
        for row in table.rows:
            values = []
            formats = []
            for raw in row:
                value, fmt = (_plain_cell(raw), None) if table.plain else _cell_value(raw)
                values.append(value)
                formats.append(fmt)
            ws.append(values)
            for cell, fmt in zip(ws[ws.max_row], formats):
                if fmt:
                    cell.number_format = fmt
                if table.plain:
                    cell.data_type = "s"  # "=..." from a Slack message is text, not a formula
            if row and str(row[0]).strip().lower() == "total":
                for cell in ws[ws.max_row]:
                    cell.font = Font(bold=True)
            widths = [max(w, len(str(r))) for w, r in zip(widths, row + [""] * len(widths))]
        for idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 10), 60)
        ws.freeze_panes = "A4"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def spreadsheet_filename(tables: List[SheetTable]) -> str:
    title = tables[0].title if tables else "hoohelp-data"
    slug = re.sub(r"[^a-z0-9]+", "-", title.lower()).strip("-")[:80] or "hoohelp-data"
    return f"{slug}.xlsx"
