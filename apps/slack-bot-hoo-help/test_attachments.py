#!/usr/bin/env python3
"""Which file a question asks for, and building xlsx / CSV / JSON from tables (no Slack)."""

import csv
import io
import json
import os
import sys
from types import SimpleNamespace

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "src"))

from attachments import build_file, file_context_line, file_format, requested_format  # noqa: E402
from spreadsheet import SheetTable  # noqa: E402

LINKS = SheetTable("Links shared in #ai (2026-10-01 to 2026-10-21)", ["Date", "Author", "URL", "Message"],
                   [["2026-10-02", "Pat Doe", "https://example.org/a", "https://uvalib.slack.com/archives/C1/p1"],
                    ["2026-10-05", "Jordan Lee", "https://example.org/b,c", "https://uvalib.slack.com/archives/C1/p2"]])


def test_file_format_from_question():
    assert file_format("can you send that as a JSON file with dates?") == "json"
    assert file_format("export it as csv please") == "csv"
    assert file_format("Give me daily gate counts for Fine Arts in Excel") == "xlsx"
    assert file_format("as a spreadsheet") == "xlsx"
    assert file_format("what are the hours?") is None
    assert file_format("excel or json, whichever") == "json"  # explicit non-Excel format wins


def test_csv_quotes_commas_and_keeps_headers():
    data, filename, title = build_file("csv", [LINKS])
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
    assert rows[0] == ["Date", "Author", "URL", "Message"]
    assert rows[2][2] == "https://example.org/b,c"
    assert filename.endswith(".csv") and filename.startswith("links-shared-in-ai")
    assert title == LINKS.title


def test_json_is_records_with_header_keys():
    data, filename, _ = build_file("json", [LINKS])
    doc = json.loads(data)
    assert doc["title"] == LINKS.title
    assert doc["rows"][0] == {"Date": "2026-10-02", "Author": "Pat Doe", "URL": "https://example.org/a",
                              "Message": "https://uvalib.slack.com/archives/C1/p1"}
    assert filename.endswith(".json")


def test_multiple_tables_json_is_a_list_and_csv_has_sections():
    other = SheetTable("Second", ["A"], [["1"]])
    doc = json.loads(build_file("json", [LINKS, other])[0])
    assert [t["title"] for t in doc] == [LINKS.title, "Second"]
    text = build_file("csv", [LINKS, other])[0].decode("utf-8-sig")
    assert LINKS.title in text and "Second" in text


def test_xlsx_still_builds():
    data, filename, _ = build_file("xlsx", [LINKS])
    assert data[:2] == b"PK" and filename.endswith(".xlsx")



def turn(role, text):
    return {"role": role, "content": [{"text": text}]}


def test_follow_up_keeps_the_format_the_thread_asked_for():  # staff trial: "all the links ever" got no CSV
    history = [turn("user", "[Doug]: all the URLs shared in #ai as a CSV file"), turn("assistant", "Attached.")]
    assert requested_format("can you give me all the links ever shared?", history) == "csv"
    assert requested_format("now as JSON please", history) == "json"
    later = history + [turn("user", "[Doug]: and as a spreadsheet"), turn("assistant", "ok")]
    assert requested_format("all of them", later) == "xlsx"
    assert requested_format("all of them", [turn("assistant", "Here is a CSV")]) is None  # only people's asks count
    assert requested_format("hours for Clemons?", []) is None


def test_file_context_line_tells_the_model_the_truth():
    assert "won't be attached" in file_context_line(None) and "CSV, JSON or Excel" in file_context_line(None)
    line = file_context_line("csv")
    assert "CSV" in line and "attached automatically" in line


MESSAGES = SheetTable("Messages in #ai (2026-10-01 to 2026-10-21)", ["Date", "Author", "Text", "Message"],
                      [["2026-10-02 09:00", "Pat Doe", '=HYPERLINK("http://evil.example","report")', "https://uvalib.slack.com/archives/C1/p1"],
                       ["2026-10-02 09:05", "+1", "**bold** and `code`", "https://uvalib.slack.com/archives/C1/p2"],
                       ["2026-10-02 09:06", "Jordan Lee", "bell\x07 here", "https://uvalib.slack.com/archives/C1/p3"]],
                      plain=True)


def test_message_text_is_never_a_formula_in_xlsx():  # review: formula injection from Slack text
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(build_file("xlsx", [MESSAGES])[0])).active
    assert ws["C4"].value == '=HYPERLINK("http://evil.example","report")' and ws["C4"].data_type == "s"
    assert ws["B5"].value == "+1" and ws["C5"].value == "**bold** and `code`"
    assert ws["C6"].value == "bell here"


def test_message_text_is_never_a_formula_in_csv():
    rows = list(csv.reader(io.StringIO(build_file("csv", [MESSAGES])[0].decode("utf-8-sig"))))
    assert rows[1][2].startswith("'=HYPERLINK") and rows[2][1] == "'+1"
    assert json.loads(build_file("json", [MESSAGES])[0])["rows"][0]["Text"].startswith("=HYPERLINK")  # JSON is data


# ---- app wiring: _attach_file (fake Slack client) ------------------------------
os.environ.setdefault("SLACK_BOT_TOKEN", "xoxb-test")
os.environ.setdefault("SLACK_SIGNING_SECRET", "test")
os.environ.setdefault("AWS_REGION", "us-east-1")
os.environ.setdefault("CONVERSATION_TRACE", "off")

GATE_STEP = {"tool_name": "OccupancyReportingRuntime___get_entrance_traffic", "status": "success",
             "output": "### Entrance Traffic Breakdown: Clemons\n\n| Entrance | In (Entries) |\n| :--- | :--- |\n| Main | 73,469 |"}


class Uploads:
    def __init__(self, error=None):
        self.files, self.posts, self.error = [], [], error

    def files_upload_v2(self, **kw):
        if self.error:
            raise RuntimeError(self.error)
        self.files.append(kw)

    def chat_postMessage(self, **kw):
        self.posts.append(kw)


class Trace:
    def __init__(self, steps):
        self.steps, self.error = steps, ""


def test_app_attaches_listing_as_json():
    import app
    client = Uploads()
    app._attach_file(client, "C9", "1.2", Trace([]), "json", [LINKS])
    assert len(client.files) == 1 and client.files[0]["filename"].endswith(".json")
    assert json.loads(client.files[0]["content"])["rows"][1]["URL"] == "https://example.org/b,c"
    assert client.files[0]["thread_ts"] == "1.2"


def test_app_csv_request_attaches_csv_not_excel():  # trial: "as a CSV" produced an .xlsx
    import app
    client = Uploads()
    app._attach_file(client, "C9", None, Trace([GATE_STEP]), "csv", [])
    assert client.files[0]["filename"].endswith(".csv")
    assert "73,469" in client.files[0]["content"].decode("utf-8-sig")


def test_app_nothing_to_attach_and_missing_scope_notice():
    import app
    client = Uploads()
    app._attach_file(client, "C9", None, Trace([]), "json", [])
    assert client.files == [] and client.posts == []
    client = Uploads(error="missing_scope")
    app._attach_file(client, "C9", None, Trace([]), "csv", [LINKS])
    assert "files:write" in client.posts[0]["text"] and "Excel" not in client.posts[0]["text"]


class FakeAgent:
    bedrock_runtime = None
    model_id = "m"

    def __init__(self):
        self.extra = None

    def process_message(self, text, *, extra_context=None, **kw):
        from session_trace import SessionTrace
        self.extra = list(extra_context or [])
        trace = SessionTrace(request_id="r1", user_message=text)
        return SimpleNamespace(text="There are 2 links in #ai.", trace=trace)

    def _sanitize_links(self, text, trace, **kw):
        return text, []


class ThreadClient(Uploads):
    def reactions_add(self, **kw):
        pass

    def reactions_remove(self, **kw):
        pass

    def users_info(self, user):
        return {"user": {"id": user, "real_name": "Doug", "profile": {}}}

    def conversations_replies(self, **kw):
        return {"messages": []}


def _follow_up(text, history, tables=(LINKS,)):
    import app
    agent, client = FakeAgent(), ThreadClient()
    tools = SimpleNamespace(listings=[SimpleNamespace(table=t) for t in tables])
    real = app.agent, app.get_bot_identity, app.slack_tools_for_event
    app.agent, app.get_bot_identity = agent, (lambda c: ("UBOT", "B1"))
    app.slack_tools_for_event = lambda *a, **k: (tools, [])
    try:
        event = {"type": "message", "channel": "C9", "channel_type": "channel", "user": "U1", "text": text,
                 "ts": "9.0", "thread_ts": "1.0"}
        app.handle_user_query(event, lambda **k: None, client, known_history=history, known_meta={})
    finally:
        app.agent, app.get_bot_identity, app.slack_tools_for_event = real
    return agent, client


def test_follow_up_in_a_csv_thread_attaches_a_csv():  # staff trial: "all the links ever" → "attached", no file
    history = [turn("user", "[Doug]: all the URLs shared in #ai as a CSV file"), turn("assistant", "Attached.")]
    agent, client = _follow_up("can you give me a listing of all the links ever shared in the channel?", history)
    assert [f["filename"].rsplit(".", 1)[1] for f in client.files] == ["csv"]
    assert any("CSV file" in line and "attached automatically" in line for line in agent.extra)


def test_retried_listing_attaches_only_the_last_one():  # default 30 days, then all-time, in one turn
    all_time = SheetTable("Links in #ai (2010-01-01 to 2026-10-22)", LINKS.headers, LINKS.rows + [LINKS.rows[0]], plain=True)
    _, client = _follow_up("all the links ever in #ai as CSV", [], tables=(LINKS, all_time))
    text = client.files[0]["content"].decode("utf-8-sig")
    assert LINKS.title not in text  # one table, no section titles
    assert client.files[0]["title"] == all_time.title and len(text.strip().splitlines()) == 4


def test_no_file_asked_tells_the_model_not_to_claim_one():
    agent, client = _follow_up("all the links in #ai", [])
    assert client.files == []
    assert any("No file will be attached" in line for line in agent.extra)


if __name__ == "__main__":
    test_file_format_from_question()
    test_csv_quotes_commas_and_keeps_headers()
    test_json_is_records_with_header_keys()
    test_multiple_tables_json_is_a_list_and_csv_has_sections()
    test_xlsx_still_builds()
    test_follow_up_keeps_the_format_the_thread_asked_for()
    test_file_context_line_tells_the_model_the_truth()
    test_message_text_is_never_a_formula_in_xlsx()
    test_message_text_is_never_a_formula_in_csv()
    test_app_attaches_listing_as_json()
    test_app_csv_request_attaches_csv_not_excel()
    test_app_nothing_to_attach_and_missing_scope_notice()
    test_follow_up_in_a_csv_thread_attaches_a_csv()
    test_retried_listing_attaches_only_the_last_one()
    test_no_file_asked_tells_the_model_not_to_claim_one()
    print("ok")
