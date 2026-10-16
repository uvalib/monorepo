#!/usr/bin/env python3
"""Unit tests for Excel attachments built from tool tables (no Slack / Bedrock)."""

import io
import os
import sys
from datetime import date

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "src"))

from openpyxl import load_workbook

from spreadsheet import (
    build_xlsx,
    spreadsheet_filename,
    tables_from_steps,
    wants_spreadsheet,
)

PERIOD_OUTPUT = """### Traffic by day: Fine Arts (2026-07-01 to 2026-07-03)
* Entries/exits counted during open hours; rows sum to the totals.

| Date | Day | Entries | Exits | Combined |
| :--- | :--- | ---: | ---: | ---: |
| 2026-07-01 | Wed | 1,200 | 1,150 | 2,350 |
| 2026-07-02 | Thu | Closed / no data | | |
| 2026-07-03 | Fri | 36 | 37 | 73 |
| Total | | 1,236 | 1,187 | 2,423 |"""

ENTRANCE_OUTPUT = """### Entrance Traffic Breakdown: Clemons
* **Date Range**: 2026-04-01 to 2026-04-30

| Entrance / Access Point | In (Entries) | Out (Exits) | Combined | % of Total |
| :--- | :--- | :--- | :--- | :--- |
| Clemons Main Entry | 17,941 | 39,540 | 57,481 | 38.4% |"""


class _Step:
    def __init__(self, tool_name, output, status="success"):
        self.tool_name = tool_name
        self.output = output
        self.status = status


def test_wants_spreadsheet():
    assert wants_spreadsheet("Can you give me daily counts for July in an Excel format?")
    assert wants_spreadsheet("export that as a spreadsheet please")
    assert wants_spreadsheet("send me a CSV")
    assert wants_spreadsheet("xlsx of Clemons gate counts")
    assert not wants_spreadsheet("What are the gate counts for Shannon this month?")
    assert not wants_spreadsheet("excellent, thanks!")


def test_tables_from_period_and_entrance_tools():
    steps = [
        _Step("OccupancyReportingRuntime___get_traffic_by_period", PERIOD_OUTPUT),
        _Step("OccupancyReportingRuntime___get_entrance_traffic", ENTRANCE_OUTPUT),
        _Step("OccupancyReportingRuntime___get_library_hours", "| a | b |\n| --- | --- |\n| 1 | 2 |"),
        _Step("OccupancyReportingRuntime___get_traffic_by_period", "Error: boom", status="error"),
    ]
    tables = tables_from_steps(steps)
    assert [t.title for t in tables] == [
        "Traffic by day: Fine Arts (2026-07-01 to 2026-07-03)",
        "Entrance Traffic Breakdown: Clemons",
    ]
    assert tables[0].headers == ["Date", "Day", "Entries", "Exits", "Combined"]
    assert len(tables[0].rows) == 4


def test_build_xlsx_types_and_sheets():
    tables = tables_from_steps(
        [
            _Step("get_traffic_by_period", PERIOD_OUTPUT),
            _Step("get_entrance_traffic", ENTRANCE_OUTPUT),
        ]
    )
    wb = load_workbook(io.BytesIO(build_xlsx(tables)))
    assert len(wb.sheetnames) == 2
    ws = wb[wb.sheetnames[0]]
    assert ws["A1"].value == "Traffic by day: Fine Arts (2026-07-01 to 2026-07-03)"
    assert [c.value for c in ws[3]] == ["Date", "Day", "Entries", "Exits", "Combined"]
    first = [c.value for c in ws[4]]
    assert first[0].date() == date(2026, 7, 1)
    assert first[2:] == [1200, 1150, 2350]
    assert [c.value for c in ws[5]][2] == "Closed / no data"
    pct = wb[wb.sheetnames[1]]["E4"]
    assert abs(pct.value - 0.384) < 1e-9
    assert pct.number_format == "0.0%"


def test_filename_from_first_table():
    tables = tables_from_steps([_Step("get_traffic_by_period", PERIOD_OUTPUT)])
    assert (
        spreadsheet_filename(tables)
        == "traffic-by-day-fine-arts-2026-07-01-to-2026-07-03.xlsx"
    )


def main():
    tests = [
        test_wants_spreadsheet,
        test_tables_from_period_and_entrance_tools,
        test_build_xlsx_types_and_sheets,
        test_filename_from_first_table,
    ]
    for fn in tests:
        fn()
        print(f"ok  {fn.__name__}")
    print(f"\n{len(tests)} tests passed")


if __name__ == "__main__":
    main()
